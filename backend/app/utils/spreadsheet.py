"""Reading uploaded .xlsx and .csv files into rows, and writing the import templates.

Nothing here knows about courses or people. It turns a file into a header row and data rows, each
with the row number the spreadsheet itself shows, and refuses a file that cannot be read as a table.
"""
import csv
import io
import re
import zipfile
from contextlib import closing
from dataclasses import dataclass
from datetime import date, datetime, time

MAX_FILE_BYTES = 5 * 1024 * 1024
# What an .xlsx may unpack to. Two thousand rows of text are well under a megabyte unpacked.
MAX_UNPACKED_BYTES = 50 * 1024 * 1024
MAX_DATA_ROWS = 2000
# Columns past this one are not read. The widest template has ten.
MAX_COLUMNS = 100
# A sheet this long is mostly empty rows left behind by formatting or a stray cell far down.
# Reading stops here instead of walking a million blank rows.
MAX_SCANNED_ROWS = 50_000
EXTENSIONS = (".xlsx", ".csv")

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MEDIA_TYPE = "text/csv; charset=utf-8"

# An old .xls workbook and a password-protected .xlsx both start with these bytes.
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"
_UNREADABLE_XLSX = "The Excel file could not be read. It may be damaged, or not a real .xlsx file."
_UNREADABLE_CSV = "The CSV file could not be read. Save it again as CSV (comma separated) and try once more."
_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d %b %Y", "%d %B %Y")
# "2024-06-15 00:00:00", as other programs write a date into a CSV.
_DATE_WITH_TIME = re.compile(r"(\d{4}-\d{1,2}-\d{1,2})[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?")
_DATE_HELP = "Date must be written like 2024-06-15 (year-month-day) or 15-06-2024 (day-month-year)"


class ImportFileError(ValueError):
    """The file as a whole cannot be imported. Nothing from it is saved."""

    status = 400


class SpreadsheetSupportMissing(ImportFileError):
    """The server lacks the package that reads and writes .xlsx. Everything else keeps working."""

    status = 503

    def __init__(self):
        super().__init__("Excel files need the openpyxl package on the server (pip install -r requirements.txt). Until it is installed, use a .csv file.")


def _openpyxl():
    # Imported on first use: a server that was updated without installing the new requirement
    # still starts and still scans faces; only Excel files are refused, with a message saying why.
    try:
        import openpyxl
    except ImportError:
        raise SpreadsheetSupportMissing() from None
    return openpyxl


@dataclass
class SheetRow:
    number: int  # the row as the spreadsheet shows it: the header of a plain file is row 1
    cells: dict  # column key -> the cell as it was read


@dataclass
class Sheet:
    headers: list  # column keys, in the order of the file
    titles: dict  # column key -> the header as it was written
    rows: list


def cell_text(value):
    """What a cell says, as text. A whole number stays whole: the number 9876543210.0 is "9876543210"."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    # Byte-order marks and zero-width spaces ride along with copied text and are never meant.
    return str(value).strip().strip("﻿​").strip()


def column_key(value):
    """How a header is recognised: "Course Code", "course-code" and " COURSE_CODE " are all course_code."""
    return re.sub(r"[^0-9a-z]+", "_", cell_text(value).lower()).strip("_")


def parse_date(value):
    """A date cell or a date written as text -> date. Raises ValueError for anything else."""
    if isinstance(value, datetime):
        found = value.date()
    elif isinstance(value, date):
        found = value
    else:
        found = None
        if isinstance(value, str):
            text = cell_text(value)
            with_time = _DATE_WITH_TIME.fullmatch(text)
            if with_time:
                text = with_time.group(1)
            for layout in _DATE_FORMATS:
                try:
                    found = datetime.strptime(text, layout).date()
                    break
                except ValueError:
                    continue
        # A bare number is not read as a date: 45458 could be an Excel serial or a typing mistake.
        if found is None:
            raise ValueError(_DATE_HELP)
    if not 1900 <= found.year <= 2100:
        raise ValueError("The year must be between 1900 and 2100")
    return found


def read_sheet(filename, content):
    """The first worksheet of an .xlsx file, or a .csv file, as a Sheet. Raises ImportFileError."""
    extension = "." + filename.rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    if extension not in EXTENSIONS:
        raise ImportFileError("Unsupported file type. Upload a .xlsx or .csv file.")
    if not content:
        raise ImportFileError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    with closing(_xlsx_rows(content) if extension == ".xlsx" else _csv_rows(content)) as rows:
        return _table(rows)


def _table(rows):
    header = None
    titles = {}
    data = []
    for number, values in rows:
        if not any(cell_text(value) for value in values):
            continue
        if header is None:
            header = []
            for index, value in enumerate(values):
                key = column_key(value)
                if not key:
                    continue
                if key in titles:
                    raise ImportFileError(f"The column \"{key}\" appears more than once in the header row.")
                titles[key] = cell_text(value)
                header.append((index, key))
            if not header:
                raise ImportFileError("The header row has no column names.")
            continue
        if len(data) >= MAX_DATA_ROWS:
            raise ImportFileError(f"The file has more than {MAX_DATA_ROWS:,} data rows. Split it into smaller files.")
        data.append(SheetRow(number, {key: values[index] if index < len(values) else None for index, key in header}))
    if header is None:
        raise ImportFileError("The file is empty.")
    return Sheet([key for _index, key in header], titles, data)


def _too_long():
    return ImportFileError(f"The sheet has more than {MAX_SCANNED_ROWS:,} rows, nearly all of them empty. Copy your data into a new sheet, or into the template, and upload that.")


def _xlsx_rows(content):
    openpyxl = _openpyxl()
    if content.startswith(_OLE_MAGIC):
        raise ImportFileError("This is an old .xls file or a password-protected workbook. Save it as a normal .xlsx file and upload that.")
    try:
        # An .xlsx is a zip. A few megabytes of it can be made to unpack into gigabytes, and the
        # text of a workbook is unpacked into memory whole, so what it would unpack to is added up
        # first, from the sizes the zip itself declares and is held to when it is read.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            unpacked = sum(item.file_size for item in archive.infolist())
    except (zipfile.BadZipFile, OSError, ValueError):
        raise ImportFileError(_UNREADABLE_XLSX) from None
    if unpacked > MAX_UNPACKED_BYTES:
        raise ImportFileError("The Excel file is too large once it is unpacked. Copy only the data rows into a new workbook, or save the sheet as CSV.")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        # Whatever the library found wrong inside the file is of no use to the person uploading it.
        raise ImportFileError(_UNREADABLE_XLSX) from None
    try:
        if not workbook.worksheets:
            raise ImportFileError("The Excel file has no worksheet.")
        sheet = workbook.worksheets[0]
        # The size a file claims for its sheet is sometimes wrong; read what is really there.
        sheet.reset_dimensions()
        for position, row in enumerate(sheet.iter_rows(max_col=MAX_COLUMNS), start=1):
            if position > MAX_SCANNED_ROWS:
                raise _too_long()
            # The row number comes from the cells themselves, so it is the one Excel shows.
            number = next((cell.row for cell in row if cell.value is not None), position)
            yield number, [cell.value for cell in row]
    except ImportFileError:
        raise
    except Exception:
        raise ImportFileError(_UNREADABLE_XLSX) from None
    finally:
        workbook.close()


def _csv_rows(content):
    text = _decode(content)
    try:
        for number, record in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=_delimiter(text)), start=1):
            if number > MAX_SCANNED_ROWS:
                raise _too_long()
            yield number, record[:MAX_COLUMNS]
    except csv.Error:
        raise ImportFileError(_UNREADABLE_CSV) from None


def _decode(content):
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings = ("utf-16",)
    elif b"\x00" in content:
        # Text never holds this byte. This is some other kind of file with its name ending in .csv.
        raise ImportFileError("This does not look like a CSV file.")
    else:
        # "CSV UTF-8" from Excel starts with a byte-order mark; plain "CSV" from Excel on Windows is cp1252.
        encodings = ("utf-8-sig", "cp1252")
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("The CSV file's text encoding is not supported. In Excel, save it as \"CSV UTF-8 (Comma delimited)\".")


def _delimiter(text):
    # Excel writes semicolons where the comma is the decimal sign. The header row shows which one is used.
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    return max(",;\t", key=lambda mark: (first_line.count(mark), mark == ","))


def build_csv(headers, rows=()):
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    # The byte-order mark is what makes Excel open the file as UTF-8.
    return buffer.getvalue().encode("utf-8-sig")


def build_template(title, columns, notes, examples):
    """An .xlsx template: an empty sheet with the header row, and a second sheet that explains it.

    columns: (name, required, what to enter) for each column, in order.
    examples: rows of made-up values, shown on the second sheet only. The first sheet holds no data,
    so a sample can never be imported by accident.
    """
    _openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    required_style = (PatternFill("solid", fgColor="0E1116"), Font(bold=True, color="FFFFFF"))
    optional_style = (PatternFill("solid", fgColor="E5E7EB"), Font(bold=True, color="0E1116"))

    def header_row(sheet, row):
        for index, (name, required, _help) in enumerate(columns, start=1):
            cell = sheet.cell(row=row, column=index, value=name)
            cell.fill, cell.font = required_style if required else optional_style
            cell.alignment = Alignment(vertical="center")

    workbook = Workbook()
    data = workbook.active
    data.title = title[:31]
    header_row(data, 1)
    data.freeze_panes = "A2"
    for index, (name, _required, _help) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        data.column_dimensions[letter].width = max(18, len(name) + 6)
        # Text format for the whole column: Excel then keeps what is typed, so an id keeps its
        # leading zeros and a phone number keeps its +91 instead of turning into 9.19877E+11.
        data.column_dimensions[letter].number_format = "@"

    guide = workbook.create_sheet("Instructions")
    guide["A1"] = f"{title} import template"
    guide["A1"].font = Font(bold=True, size=14)
    row = 3
    guide.cell(row=row, column=1, value="How to use").font = Font(bold=True)
    for note in notes:
        row += 1
        guide.cell(row=row, column=1, value=f"• {note}")
    row += 2
    for index, heading in enumerate(("Column", "Required", "What to enter"), start=1):
        guide.cell(row=row, column=index, value=heading).font = Font(bold=True)
    for name, required, help_text in columns:
        row += 1
        guide.cell(row=row, column=1, value=name)
        guide.cell(row=row, column=2, value="Yes" if required else "No")
        guide.cell(row=row, column=3, value=help_text)
    row += 2
    guide.cell(row=row, column=1, value="Example rows (made-up values, for illustration only)").font = Font(bold=True)
    row += 1
    header_row(guide, row)
    for example in examples:
        row += 1
        for index, value in enumerate(example, start=1):
            cell = guide.cell(row=row, column=index, value=value)
            cell.number_format = "@"
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 26
    guide.column_dimensions["C"].width = 70
    for index in range(4, len(columns) + 1):
        guide.column_dimensions[get_column_letter(index)].width = 22

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
