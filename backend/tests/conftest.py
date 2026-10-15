import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import os
import secrets

# Set, not defaulted: whatever DATABASE_URL this shell or backend/.env holds, no test can reach it.
# Every test runs on an in-memory SQLite database of its own.
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
os.environ["FACE_SCAN_LOG"] = "false"
os.environ["FACE_TIMING_LOG"] = "false"
# The tests fix their own settings, whatever backend/.env says on this machine. Enrollment tests
# use the full five poses unless a test asks for a quick enrollment.
os.environ.update(FACE_ENROLL_MODE="guided", FACE_ENROLL_QUICK_STEPS="front,left,right", FACE_ENROLL_BURST_FRAMES="3")
# A signing secret made for this test run and gone with it. The one in backend/.env is never read.
os.environ["JWT_SECRET"] = secrets.token_urlsafe(48)
# The tests call the backend directly, not through the frontend's /api proxy, so the cookie's path
# is the backend's own. The rest are the documented defaults.
os.environ.update(AUTH_COOKIE_PATH="/auth", AUTH_ACCESS_MINUTES="15", AUTH_REFRESH_DAYS="7", API_DOCS="false", FRONTEND_ORIGIN="http://localhost:5173")

from datetime import timedelta

import pytest
from argon2 import PasswordHasher
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models.admin_account import AdminAccount, AdminSession
from app.models.audit_log import AuditLog  # noqa: F401  (its table is created with the others)
from app.services import audit_service, auth_service, password_service  # noqa: F401  (audit_service: the face enrollment listener)

PRODUCTION_HASHER = password_service._hasher
# The same algorithm with the smallest cost it allows. Hundreds of tests create accounts; what
# they prove does not depend on a hash taking a tenth of a second.
FAST_HASHER = PasswordHasher(time_cost=1, memory_cost=8, parallelism=1)


@pytest.fixture(autouse=True)
def auth_memory(monkeypatch):
    # Every test starts with cheap hashing and with no attempts counted and no login remembered.
    monkeypatch.setattr(password_service, "_hasher", FAST_HASHER)
    auth_service.reset_memory()
    auth_service._decoy_hash.cache_clear()
    yield
    auth_service.reset_memory()
    auth_service._decoy_hash.cache_clear()


@pytest.fixture()
def production_hashing(monkeypatch):
    # For the few tests that are about the real cost settings.
    monkeypatch.setattr(password_service, "_hasher", PRODUCTION_HASHER)


def add_account(db, email="admin@example.com", password=None, role="admin", **fields):
    """A login account, saved. Without a password it gets a hash that no password matches."""
    account = AdminAccount(
        email=email, full_name=fields.pop("full_name", "Sample Admin"), role=role,
        password_hash=password_service.hash_password(password) if password else "no-password-set", **fields,
    )
    db.add(account)
    db.commit()
    db.refresh(account)
    return account


def open_login(db, account):
    """Logs an account in without going through the login route. Returns (access token, refresh token)."""
    refresh_token = secrets.token_urlsafe(32)
    now = auth_service.clock()
    session = AdminSession(account_id=account.id, token_hash=auth_service.hash_token(refresh_token), expires_at=now + timedelta(days=7), last_used_at=now)
    db.add(session)
    db.commit()
    access_token, _seconds = auth_service.issue_access_token(account, session, now)
    return access_token, refresh_token


@pytest.fixture()
def make_account():
    return add_account


@pytest.fixture()
def log_in():
    return open_login


@pytest.fixture()
def db_session():
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    session_factory = sessionmaker(bind=engine)
    session = session_factory()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def backend(monkeypatch):
    # The real application over HTTP, on its own empty in-memory database. Sessions are made the
    # way app/database.py makes them (no autoflush) and handed out by the real get_db, so a route
    # behaves here as it does when the backend runs. Calling backend() gives a new browser, so to
    # speak: a client with its own cookies and nobody logged in. backend.db() opens a session on
    # the same database, for a test to prepare rows or to look at what a request left behind.
    from fastapi.testclient import TestClient
    from sqlalchemy.pool import StaticPool

    from app import database
    from app.main import app

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=engine)
    session_factory = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    monkeypatch.setattr(database, "SessionLocal", session_factory)

    def connect(**options):
        # Not entered as a context manager: the startup hook (the schema of the real database,
        # loading the face models) must not run for these tests. https, because the refresh cookie
        # is marked Secure and no client sends such a cookie over plain http. The header is the one
        # the console sends with its auth requests (app/routes/auth_routes.py, CLIENT_HEADER).
        headers = {"X-Requested-With": "vernex", **options.pop("headers", {})}
        test_client = TestClient(app, base_url="https://testserver", headers=headers, **options)
        test_client.db = session_factory
        return test_client

    connect.db = session_factory
    try:
        yield connect
    finally:
        engine.dispose()


@pytest.fixture()
def anon_client(backend):
    # Nobody is logged in.
    return backend()


@pytest.fixture()
def client(backend):
    # Logged in as an admin, which is what every route except the public few now asks for.
    # client.db() opens a session on the same database.
    test_client = backend()
    with backend.db() as db:
        account = add_account(db, "admin@example.com")
        access_token, _refresh_token = open_login(db, account)
    test_client.headers["Authorization"] = f"Bearer {access_token}"
    test_client.account_email = "admin@example.com"
    return test_client


@pytest.fixture()
def make_xlsx():
    # rows: one list of cell values per spreadsheet row, starting at row 1. None leaves a row empty,
    # as a row someone skipped in Excel is.
    def build(rows, title="Sheet1", more_sheets=()):
        import io

        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = title
        for number, values in enumerate(rows, start=1):
            for column, value in enumerate(values or [], start=1):
                sheet.cell(row=number, column=column, value=value)
        for name, other_rows in more_sheets:
            other = workbook.create_sheet(name)
            for values in other_rows:
                other.append(values)
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    return build


@pytest.fixture()
def make_csv():
    # rows: one list of cell values per line. None is an empty line.
    def build(rows, encoding="utf-8", delimiter=","):
        import csv
        import io

        buffer = io.StringIO(newline="")
        writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\r\n")
        for values in rows:
            writer.writerow(values or [])
        return buffer.getvalue().encode(encoding)

    return build


class Clock:
    # The time the face routes see. Every look at it moves it on a little, like real frames do;
    # a test can also jump it forward.
    def __init__(self, step=0.3):
        self.now = 1000.0
        self.step = step

    def __call__(self):
        self.now += self.step
        return self.now

    def advance(self, seconds):
        self.now += seconds


@pytest.fixture()
def clock(monkeypatch):
    from app.routes import face_routes
    from app.services import tracking_service

    clock = Clock()
    monkeypatch.setattr(face_routes, "clock", clock)
    face_routes.pending_enrollments.clear()
    face_routes.gallery_cache.clear()
    face_routes.session_cache.clear()
    face_routes.marked_today.clear()
    tracking_service.reset()
    yield clock
    face_routes.pending_enrollments.clear()
    face_routes.gallery_cache.clear()
    face_routes.session_cache.clear()
    face_routes.marked_today.clear()
    tracking_service.reset()
