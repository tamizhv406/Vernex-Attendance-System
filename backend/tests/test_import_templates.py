import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.models.course import Course, CourseBatch
from app.models.person import Person
from app.services.import_service import SPECS
from app.utils.spreadsheet import XLSX_MEDIA_TYPE, read_sheet

# The upload dialog's column lists are tested against this same file (frontend/src/lib/imports.test.js).
COLUMNS = json.loads((Path(__file__).parent / "data" / "import_columns.json").read_text(encoding="utf-8"))
# kind of record -> (template address, upload address)
ADDRESSES = {
    "courses": ("/courses/upload/template", "/courses/upload"),
    "batches": ("/batches/upload/template", "/batches/upload"),
    "students": ("/users/upload/template?person_type=student", "/users/upload?person_type=student"),
    "staff": ("/users/upload/template?person_type=admin", "/users/upload?person_type=admin"),
    "members": ("/users/upload/template?person_type=member", "/users/upload?person_type=member"),
}
KINDS = list(ADDRESSES)


def workbook_of(response):
    return load_workbook(io.BytesIO(response.content))


def example_rows(workbook):
    # On the Instructions sheet, under the line that announces them: a header row, then the examples.
    rows = [[cell.value for cell in row] for row in workbook["Instructions"].iter_rows()]
    start = next(index for index, row in enumerate(rows) if str(row[0] or "").startswith("Example rows"))
    return rows[start + 1], [row for row in rows[start + 2:] if any(row)]


def create_what_the_examples_refer_to(client, kind):
    # The examples for batches name the course CS101; those for people name CS101 and its batch
    # "Morning A". Neither is created by the file itself, so both must be there first.
    if kind == "courses":
        return
    course = client.post("/courses/create", json={"course_code": "CS101", "course_name": "Computer Science"}).json()["data"]["id"]
    if kind != "batches":
        client.post("/batches/create", json={"course_id": course, "batch_name": "Morning A"})


def test_the_columns_of_every_upload_are_the_documented_ones():
    assert sorted(SPECS) == sorted(key for key in COLUMNS if key != "about")
    for kind, spec in SPECS.items():
        assert spec.required == COLUMNS[kind]["required"], kind
        assert spec.optional == COLUMNS[kind]["optional"], kind
        # Required columns first, then the optional ones: the order of the template.
        assert [column.name for column in spec.columns] == COLUMNS[kind]["required"] + COLUMNS[kind]["optional"], kind


@pytest.mark.parametrize("kind", KINDS)
def test_a_template_is_a_workbook_with_the_header_row_and_nothing_to_import(client, kind):
    response = client.get(ADDRESSES[kind][0])

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["content-disposition"] == f'attachment; filename="{kind}_import_template.xlsx"'
    workbook = workbook_of(response)
    assert workbook.sheetnames == [SPECS[kind].title, "Instructions"]
    data = workbook.worksheets[0]
    assert [cell.value for cell in data[1]] == COLUMNS[kind]["required"] + COLUMNS[kind]["optional"]
    # Only the header row: a sample can never be imported by accident.
    assert data.max_row == 1
    # Whatever is typed into a column is kept as typed (an id's leading zeros, a number's +91).
    assert {data.column_dimensions[letter].number_format for letter in "ABC"} == {"@"}


@pytest.mark.parametrize("kind", KINDS)
def test_a_template_explains_its_columns_and_shows_made_up_examples(client, kind):
    workbook = workbook_of(client.get(ADDRESSES[kind][0]))
    guide = [[cell.value for cell in row if cell.value is not None] for row in workbook["Instructions"].iter_rows()]

    explained = {row[0]: row[1] for row in guide if len(row) == 3 and row[1] in ("Yes", "No")}
    assert explained == {**{name: "Yes" for name in COLUMNS[kind]["required"]}, **{name: "No" for name in COLUMNS[kind]["optional"]}}
    header, examples = example_rows(workbook)
    assert [name for name in header if name] == COLUMNS[kind]["required"] + COLUMNS[kind]["optional"]
    assert 1 <= len(examples) <= 2
    # Nobody's real details: the names say what they are.
    if "full_name" in header:
        assert all(str(row[header.index("full_name")]).startswith("Sample ") for row in examples)


@pytest.mark.parametrize("kind", KINDS)
def test_an_untouched_template_is_refused_because_it_has_no_rows(client, kind):
    template = client.get(ADDRESSES[kind][0]).content

    response = client.post(ADDRESSES[kind][1], files={"file": (f"{kind}.xlsx", template)})

    assert response.status_code == 400
    assert response.json() == {"detail": "The file has a header row but no data rows."}
    with client.db() as db:
        assert (db.query(Course).count(), db.query(CourseBatch).count(), db.query(Person).count()) == (0, 0, 0)


@pytest.mark.parametrize("kind", KINDS)
def test_a_template_filled_in_with_its_own_examples_imports_without_renaming_anything(client, kind):
    create_what_the_examples_refer_to(client, kind)
    workbook = workbook_of(client.get(ADDRESSES[kind][0]))
    _header, examples = example_rows(workbook)
    data = workbook.worksheets[0]
    for example in examples:
        data.append(example[:data.max_column])
    filled = io.BytesIO()
    workbook.save(filled)

    response = client.post(ADDRESSES[kind][1], files={"file": (f"{kind}.xlsx", filled.getvalue())})

    assert response.status_code == 200, response.text
    assert response.json()["data"]["errors"] == []
    assert response.json()["data"]["imported"] == len(examples)
    assert response.json()["data"]["ignored_columns"] == []


@pytest.mark.parametrize("kind", KINDS)
def test_the_csv_template_is_the_header_row(client, kind):
    template_address = ADDRESSES[kind][0]
    response = client.get(template_address + ("&" if "?" in template_address else "?") + "format=csv")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"] == f'attachment; filename="{kind}_import_template.csv"'
    sheet = read_sheet("template.csv", response.content)
    assert sheet.headers == COLUMNS[kind]["required"] + COLUMNS[kind]["optional"]
    assert sheet.rows == []


def test_staff_templates_can_be_asked_for_by_the_word_staff(client):
    assert client.get("/users/upload/template?person_type=staff").headers["content-disposition"] == 'attachment; filename="staff_import_template.xlsx"'


@pytest.mark.parametrize("address, said", [
    ("/users/upload/template", "person_type must be student, admin or member"),
    ("/users/upload/template?person_type=teacher", "person_type must be student, admin or member"),
    ("/courses/upload/template?format=pdf", "The template is available as xlsx or csv."),
])
def test_a_template_that_does_not_exist_is_refused_plainly(client, address, said):
    response = client.get(address)

    assert response.status_code == 400
    assert response.json() == {"detail": said}
